"""
app.py

Interface Streamlit du moteur de recherche documentaire (projet 3 + étapes
recherche par sens, RAG / projet 4, et analyse de clauses / projet 5).

Reprend l'import de documents du projet 2 (fichiers.py, extraction.py,
database.py) et propose trois façons d'interroger le CONTENU TEXTUEL des
documents (pas seulement leurs métadonnées comme au projet 2) :
- par MOTS-CLÉS, grâce à l'index plein texte FTS5 (projet 3) ;
- par SENS, grâce aux embeddings OpenAI (étape embeddings) : retrouve un
  passage même s'il ne partage aucun mot avec la requête, tant que le sens
  est proche ;
- en POSANT UNE QUESTION (projet 4, RAG) : retrouve les passages pertinents
  par sens, puis les envoie à un LLM qui rédige une vraie réponse -- tout en
  affichant séparément les extraits bruts, pour rester vérifiable.

S'y ajoute une section "Analyser une clause" (projet 5) : pour un document
et un type de clause tapé librement, détecte si la clause existe (même
principe que "par sens" et le RAG), la fait analyser par un LLM avec un
niveau de risque, et permet d'exporter l'historique accumulé en Excel.

Enfin, une section "Analyse du data room" (projet final) généralise ce
principe à TOUS les documents x une liste fixe de 24 catégories juridiques
standards, en parallèle (due_diligence.py), avec une matrice de risques
croisée et une synthèse rédigée par le LLM à partir d'une agrégation SQL.
"""

import io
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import clauses
import database
import due_diligence
import embeddings
import extraction
import fichiers
import llm

CATEGORIES = [
    "Corporate", "Contrats", "Contentieux", "Fiscal", "Social",
    "Immobilier", "Financement", "Propriété intellectuelle",
    "Réglementaire", "Autre",
]

FORMATS_ACCEPTES = ["pdf", "docx", "txt"]

database.initialiser_base()

st.title("Moteur de recherche documentaire")

# ----------------------------------------------------------------------
# Import d'un document
# ----------------------------------------------------------------------
st.header("Importer un document")

fichier_uploade = st.file_uploader(
    "Choisir un fichier (PDF, DOCX ou TXT)",
    type=FORMATS_ACCEPTES,
)

categorie_import = st.selectbox("Catégorie juridique", CATEGORIES)
commentaire_import = st.text_input("Commentaire (optionnel)")

if st.button("Importer le document"):
    if fichier_uploade is None:
        st.warning("Choisis d'abord un fichier.")
    else:
        contenu = fichier_uploade.getvalue()
        if len(contenu) == 0:
            st.error("Ce fichier est vide, il n'a pas été importé.")
        else:
            type_fichier = Path(fichier_uploade.name).suffix.lower().lstrip(".")
            chemin = fichiers.sauvegarder_fichier(fichier_uploade.name, contenu)
            commentaire_final = commentaire_import if commentaire_import else None

            id_document = database.ajouter_document(
                nom=fichier_uploade.name,
                type_fichier=type_fichier,
                categorie=categorie_import,
                chemin=chemin,
                commentaire=commentaire_final,
            )

            # L'extraction alimente l'index de recherche par mots-clés (FTS5).
            # Une erreur ici n'empêche pas l'import : le document existe et
            # reste consultable, il ne sera simplement pas trouvable par une
            # recherche de contenu.
            try:
                texte = extraction.extraire_texte(chemin, type_fichier)
                database.indexer_document(id_document, fichier_uploade.name, texte)
                if texte.strip() == "":
                    st.warning(
                        f"Document '{fichier_uploade.name}' importé, mais aucun texte n'a pu "
                        "en être extrait (PDF scanné en image ?) : il ne sera pas trouvable "
                        "par une recherche de contenu."
                    )
                else:
                    st.success(f"Document '{fichier_uploade.name}' importé et indexé avec succès.")

                    # Indexation pour la recherche par SENS (appel à l'API
                    # OpenAI). Séparée de l'indexation FTS5 ci-dessus : si
                    # elle échoue (pas de clé API, pas de réseau, quota
                    # dépassé...), le document reste importé et trouvable
                    # par mots-clés, il ne sera simplement pas trouvable par
                    # une recherche par sens.
                    try:
                        nb_chunks = database.indexer_embeddings(id_document, texte)
                        st.info(f"{nb_chunks} passage(s) indexé(s) pour la recherche par sens.")
                    except embeddings.CleApiManquante:
                        st.info(
                            "Recherche par sens non disponible pour ce document (aucune clé "
                            "API OpenAI configurée dans .env) : il reste trouvable par mots-clés."
                        )
                    except Exception as erreur:
                        st.info(
                            f"Recherche par sens non disponible pour ce document ({erreur}) : "
                            "il reste trouvable par mots-clés."
                        )
            except (extraction.FormatNonSupporte, extraction.ErreurExtraction) as erreur:
                database.indexer_document(id_document, fichier_uploade.name, "")
                st.warning(f"Document importé, mais non indexé : {erreur}")

            st.rerun()

st.divider()

# ----------------------------------------------------------------------
# Recherche dans le contenu des documents, trois modes possibles. Les
# trois interrogent le même contenu texte, mais pas de la même façon :
# mots-clés compare des mots, par sens compare des vecteurs numériques
# représentant le sens, et poser une question ajoute une étape de
# génération : le LLM rédige une réponse à partir des passages retrouvés
# par sens (RAG -- retrieval-augmented generation, projet 4).
# ----------------------------------------------------------------------
st.header("Rechercher dans les documents")

mode_recherche = st.radio(
    "Mode de recherche",
    ["Mots-clés (FTS5)", "Par sens (embeddings)", "Poser une question (IA)"],
    horizontal=True,
    help=(
        "Mots-clés : retrouve les mots tapés tels quels (ou leur préfixe). "
        "Par sens : retrouve les passages dont le SENS se rapproche de la "
        "requête, même avec des mots différents. Poser une question : en "
        "plus de retrouver les passages, un LLM rédige une réponse à "
        "partir d'eux. Les deux derniers modes nécessitent une clé API "
        "OpenAI configurée dans .env."
    ),
)
recherche_texte = st.text_input(
    "Rechercher dans le contenu, ou poser une question",
    placeholder="ex. clause de non-concurrence, ou : quel est le préavis de résiliation du bail ?",
)
categorie_filtre = st.selectbox("Filtrer par catégorie", ["Toutes"] + CATEGORIES)

if mode_recherche == "Mots-clés (FTS5)":
    resultats = database.rechercher(recherche_texte, categorie_filtre)

    st.write(f"{len(resultats)} document(s) trouvé(s)")

    # ------------------------------------------------------------------
    # Affichage des résultats, aperçu, modification, suppression
    # ------------------------------------------------------------------
    for resultat in resultats:
        id_doc, nom, type_fichier, categorie, date_ajout, chemin, commentaire, extrait = resultat

        with st.expander(f"{nom} — {categorie} — {date_ajout}"):
            st.write(f"Type : {type_fichier}")
            st.write(f"Chemin : {chemin}")

            if extrait:
                st.markdown(f"Extrait correspondant : {extrait}")

            if st.button("Voir un aperçu du texte", key=f"apercu_{id_doc}"):
                try:
                    texte = extraction.extraire_texte(chemin, type_fichier)
                    if texte.strip() == "":
                        st.warning("Aucun texte n'a pu être extrait de ce document.")
                    else:
                        st.text_area("Aperçu", texte[:1000], height=200, key=f"texte_{id_doc}")
                except extraction.FormatNonSupporte:
                    st.info("L'aperçu n'est pas disponible pour ce format.")
                except extraction.ErreurExtraction as erreur:
                    st.error(str(erreur))

            st.subheader("Modifier")
            index_categorie = CATEGORIES.index(categorie) if categorie in CATEGORIES else 0
            nouvelle_categorie = st.selectbox(
                "Catégorie", CATEGORIES, index=index_categorie, key=f"cat_{id_doc}",
            )
            nouveau_commentaire = st.text_input(
                "Commentaire", value=commentaire or "", key=f"com_{id_doc}",
            )
            if st.button("Enregistrer les modifications", key=f"save_{id_doc}"):
                database.modifier_document(id_doc, nouvelle_categorie, nouveau_commentaire or None)
                st.success("Document mis à jour.")
                st.rerun()

            st.subheader("Supprimer")
            confirmer = st.checkbox("Confirmer la suppression", key=f"confirm_{id_doc}")
            if st.button("Supprimer ce document", key=f"delete_{id_doc}"):
                if confirmer:
                    database.supprimer_document(id_doc)
                    fichiers.supprimer_fichier(chemin)
                    st.success("Document supprimé.")
                    st.rerun()
                else:
                    st.warning("Coche la case de confirmation avant de supprimer.")

elif mode_recherche == "Par sens (embeddings)":
    # Un appel API est nécessaire pour calculer l'embedding de la requête
    # elle-même, donc les mêmes erreurs qu'à l'import peuvent survenir (clé
    # absente, réseau, quota). On les affiche clairement au lieu de faire
    # planter la page.
    if not recherche_texte.strip():
        resultats_sens = []
    else:
        try:
            resultats_sens = database.rechercher_par_sens(recherche_texte, categorie_filtre)
        except embeddings.CleApiManquante as erreur:
            st.error(str(erreur))
            resultats_sens = []
        except Exception as erreur:
            st.error(f"Recherche par sens indisponible : {erreur}")
            resultats_sens = []

    st.write(f"{len(resultats_sens)} passage(s) trouvé(s)")

    for resultat in resultats_sens:
        id_doc, nom, type_fichier, categorie, date_ajout, chemin, commentaire, texte_chunk, score = resultat

        with st.expander(f"{nom} — {categorie} — proximité de sens {score:.2f}"):
            st.write(f"Type : {type_fichier}")
            st.write(f"Chemin : {chemin}")
            st.markdown(f"Passage correspondant : *{texte_chunk}*")

else:
    # Mode "Poser une question" (RAG, projet 4) : on retrouve d'abord les
    # passages pertinents par sens (comme le mode précédent), puis on les
    # envoie à un LLM qui rédige une réponse à partir d'eux -- jamais à
    # partir de sa propre mémoire des documents, puisqu'il n'en a aucune.
    if not recherche_texte.strip():
        chunks_pertinents = []
    else:
        try:
            resultats_bruts = database.rechercher_par_sens(
                recherche_texte, categorie_filtre, top_n=llm.NB_CHUNKS_CONTEXTE,
            )
            chunks_pertinents = llm.filtrer_chunks_pertinents(resultats_bruts)
        except embeddings.CleApiManquante as erreur:
            st.error(str(erreur))
            chunks_pertinents = []
        except Exception as erreur:
            st.error(f"Recherche indisponible : {erreur}")
            chunks_pertinents = []

    if recherche_texte.strip() and not chunks_pertinents:
        st.info("Aucun passage suffisamment pertinent n'a été trouvé pour répondre à cette question.")

    if chunks_pertinents:
        # La génération de la réponse est un second appel API, séparé de la
        # recherche des passages ci-dessus : s'il échoue (clé absente,
        # quota, réseau...), les passages bruts restent affichés plus bas,
        # seule la synthèse rédigée est indisponible.
        try:
            reponse = llm.generer_reponse(recherche_texte, chunks_pertinents)

            st.success(f"Réponse trouvée dans le document : {reponse['reponse_document']}")

            if reponse["complement_connaissance_generale"]:
                st.warning(
                    "Connaissance générale du LLM (hors document, à vérifier) : "
                    f"{reponse['complement_connaissance_generale']}"
                )
        except embeddings.CleApiManquante as erreur:
            st.error(str(erreur))
        except Exception as erreur:
            st.error(f"Impossible de générer une réponse rédigée : {erreur}")

        # Sources : les extraits bruts tels que stockés en base, jamais
        # reformulés par le LLM -- c'est sur eux que la réponse peut être
        # vérifiée.
        st.subheader("Sources")
        for resultat in chunks_pertinents:
            id_doc, nom, type_fichier, categorie, date_ajout, chemin, commentaire, texte_chunk, score = resultat
            with st.expander(f"{nom} — {categorie} — proximité de sens {score:.2f}"):
                st.write(f"Type : {type_fichier}")
                st.write(f"Chemin : {chemin}")
                st.markdown(f"Extrait original : *{texte_chunk}*")

st.divider()

# ----------------------------------------------------------------------
# Analyse de clauses (projet 5) : pour un document et un type de clause
# tapé librement, détecte si la clause existe (même principe que "par
# sens" ci-dessus : au moins un chunk au-dessus du seuil de pertinence),
# la fait analyser par un LLM avec un niveau de risque, et accumule
# l'historique en base pour un export Excel global.
# ----------------------------------------------------------------------
st.header("Analyser une clause")

documents_existants = database.lire_documents()

if not documents_existants:
    st.info("Importe d'abord un document pour pouvoir y analyser une clause.")
else:
    noms_documents = {f"{doc[1]} (id {doc[0]})": doc[0] for doc in documents_existants}
    nom_document_choisi = st.selectbox("Document à analyser", list(noms_documents.keys()))
    doc_id_choisi = noms_documents[nom_document_choisi]

    type_clause = st.text_input(
        "Type de clause à rechercher",
        placeholder="ex. non-concurrence, confidentialité, résiliation",
    )

    if st.button("Analyser cette clause"):
        if not type_clause.strip():
            st.warning("Tape d'abord un type de clause.")
        else:
            # Toute la logique de détection en deux étapes (candidats par
            # similarité de sens, puis vérification + analyse par le LLM)
            # est dans clauses.detecter_et_analyser() -- partagée avec
            # due_diligence.py, qui l'utilise en parallèle sur tous les
            # documents x les 24 catégories fixes du data room. Ici, on ne
            # fait qu'afficher le message correspondant au statut retourné.
            resultat = clauses.detecter_et_analyser(doc_id_choisi, type_clause)

            if resultat["statut"] == "detection_indisponible":
                st.error(f"Détection indisponible : {resultat['erreur']}")
            else:
                if resultat["statut"] == "absente":
                    st.info(f"Aucune clause de type « {type_clause} » détectée dans ce document.")
                elif resultat["statut"] == "presente":
                    st.success("Clause détectée et analysée.")
                elif resultat["statut"] == "detectee_non_analysee":
                    st.warning(f"Clause candidate détectée mais non vérifiée/analysée : {resultat['erreur']}")

                database.enregistrer_analyse_clause(
                    resultat["doc_id"], resultat["type_clause"], resultat["clause_presente"],
                    resultat["texte_extrait"], resultat["analyse"], resultat["niveau_risque"],
                )
                st.rerun()

    # ------------------------------------------------------------------
    # Historique des analyses accumulées (tous documents confondus) et
    # export Excel, sur le principe du projet 1 (pandas + openpyxl, en
    # mémoire via BytesIO, pas de fichier temporaire sur disque).
    # ------------------------------------------------------------------
    st.subheader("Historique des analyses")

    BADGES_RISQUE = {"élevé": st.error, "moyen": st.warning, "faible": st.success}

    analyses = database.lire_analyses_clauses()

    if not analyses:
        st.write("Aucune analyse enregistrée pour l'instant.")
    else:
        for ligne in analyses:
            id_analyse, nom, type_clause_ligne, presente, texte_extrait, analyse_texte, niveau_risque, date_analyse = ligne

            titre = f"{nom} — {type_clause_ligne} — {'présente' if presente else 'absente'}"
            with st.expander(titre):
                st.write(f"Date de l'analyse : {date_analyse}")
                if not presente:
                    st.write("Clause non détectée dans ce document.")
                else:
                    st.markdown(f"Extrait original : *{texte_extrait}*")
                    if analyse_texte and niveau_risque:
                        afficher_badge = BADGES_RISQUE.get(niveau_risque, st.write)
                        afficher_badge(f"Niveau de risque : {niveau_risque}")
                        st.write(analyse_texte)
                    else:
                        st.info("Clause détectée mais pas encore analysée (l'analyse par IA a échoué).")

        # Construction du fichier Excel en mémoire (io.BytesIO) : aucun
        # fichier n'est écrit sur le disque, il n'existe que le temps du
        # téléchargement par l'utilisateur.
        colonnes = [
            "Document", "Type de clause", "Présente", "Extrait", "Analyse", "Niveau de risque", "Date",
        ]
        lignes_export = [
            (nom, type_clause_ligne, "Oui" if presente else "Non", texte_extrait, analyse_texte, niveau_risque, date_analyse)
            for (_, nom, type_clause_ligne, presente, texte_extrait, analyse_texte, niveau_risque, date_analyse) in analyses
        ]
        df_export = pd.DataFrame(lignes_export, columns=colonnes)

        fichier_excel = io.BytesIO()
        with pd.ExcelWriter(fichier_excel, engine="openpyxl") as writer:
            df_export.to_excel(writer, index=False, sheet_name="Analyses de clauses")
            feuille = writer.sheets["Analyses de clauses"]

            for cellule in feuille[1]:
                cellule.font = Font(bold=True)
                cellule.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
                cellule.alignment = Alignment(horizontal="center")

            for colonne in feuille.columns:
                largeur_max = max((len(str(cellule.value)) for cellule in colonne if cellule.value is not None), default=0)
                lettre_colonne = get_column_letter(colonne[0].column)
                feuille.column_dimensions[lettre_colonne].width = min(largeur_max + 2, 60)

        fichier_excel.seek(0)

        st.download_button(
            label="Télécharger l'historique des analyses en Excel",
            data=fichier_excel.getvalue(),
            file_name="analyses_clauses.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

st.divider()

# ----------------------------------------------------------------------
# Analyse du data room (projet final) : généralise "Analyser une clause"
# ci-dessus à TOUS les documents x une liste fixe de 24 catégories
# juridiques standards (due_diligence.py), exécutées en parallèle. Produit
# une matrice de risques croisée et une synthèse rédigée par le LLM à
# partir d'une agrégation SQL (jamais des analyses brutes).
# ----------------------------------------------------------------------
st.header("Analyse du data room")

st.write(
    "Analyse tous les documents importés sur une liste fixe de 24 catégories "
    "juridiques standards de due diligence (parties, durée, résiliation, "
    "propriété intellectuelle, confidentialité, données personnelles...), "
    "en parallèle. Peut prendre plusieurs minutes selon le nombre de documents."
)

if not documents_existants:
    st.info("Importe d'abord des documents pour lancer l'analyse du data room.")
else:
    if st.button("Lancer l'analyse du data room"):
        barre_progression = st.progress(0.0)
        texte_progression = st.empty()

        def afficher_progression(nb_fait, nb_total):
            barre_progression.progress(nb_fait / nb_total)
            texte_progression.write(f"{nb_fait} / {nb_total} cases analysées (document x catégorie)")

        with st.spinner("Analyse du data room en cours..."):
            due_diligence.lancer_analyse_data_room(on_resultat=afficher_progression)

        st.success("Analyse du data room terminée.")
        st.rerun()

    st.subheader("Matrice de risques")

    matrice = due_diligence.construire_matrice(database.lire_analyses_clauses())

    if matrice.empty:
        st.write("Aucune analyse du data room enregistrée pour l'instant.")
    else:
        # Coloration des cellules selon le niveau de risque -- même esprit
        # que les badges de la section "Analyser une clause" ci-dessus,
        # mais appliqué à toute la matrice d'un coup.
        COULEURS_RISQUE = {
            "Élevé": "background-color: #f8d7da",
            "Moyen": "background-color: #fff3cd",
            "Faible": "background-color: #d4edda",
            "Absente": "background-color: #f0f0f0",
            "Détectée (non analysée)": "background-color: #e2e3e5",
        }
        st.dataframe(matrice.style.map(lambda valeur: COULEURS_RISQUE.get(valeur, "")))

        st.subheader("Synthèse")
        if st.button("Générer la synthèse du data room"):
            # Agrégation SQL d'abord (comptage fiable, fait par le code),
            # puis un seul appel LLM séparé pour rédiger le texte -- jamais
            # les analyses brutes envoyées au modèle.
            try:
                comptages_texte = due_diligence.agreger_resultats_data_room()
                synthese = due_diligence.generer_synthese(comptages_texte)
                st.write(synthese)
            except embeddings.CleApiManquante as erreur:
                st.error(str(erreur))
            except Exception as erreur:
                st.error(f"Impossible de générer la synthèse : {erreur}")
